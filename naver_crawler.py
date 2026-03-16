from datetime import datetime, timedelta
import html
import json
import os
import time
from urllib.parse import quote

import numpy as np
import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from selenium import webdriver
from selenium.webdriver.chrome.options import Options


# Crawling configuration
KEYWORDS = ["이재명", "김문수", "이준석", "대선"]
START_DATE = "2025.05.19"
END_DATE = "2025.05.21"
MAX_ARTICLES_PER_KEYWORD = 6000
HEADLESS = True
SEARCH_PAGE_START = 1
SEARCH_PAGE_END = 2000
SEARCH_PAGE_STEP = 10
MAX_CONSECUTIVE_EMPTY_PAGES = 5

# Output configuration
ARTICLES_XLSX = "naver_news_articles.xlsx"
COMMENTS_RAW_CSV = "naver_comments_raw.csv"
COMMENTS_FILTERED_CSV = "naver_comments_filtered.csv"
BOT_STATS_CSV = "naver_comment_bot_stats.csv"

HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}


def _str_to_date(value):
    return datetime.strptime(value, "%Y.%m.%d")


def _build_driver(headless=True):
    options = Options()
    if headless:
        options.add_argument("--headless")
    options.add_argument("--disable-gpu")
    options.add_argument("window-size=1920x1080")
    return webdriver.Chrome(options=options)


def _parse_article_title_and_date(soup):
    title = (
        soup.select_one(".media_end_head_headline")
        or soup.find("h2")
        or soup.find("div", class_="news_title")
        or soup.find("div", class_="end_tit")
    )
    date = (
        soup.select_one("span.media_end_head_info_datestamp_time._ARTICLE_DATE_TIME")
        or soup.find("em", class_="date")
        or soup.find("span", class_="date")
        or soup.find("div", class_="date")
    )
    if title and date:
        return title.text.strip(), date.text.strip()
    return None, None


def fetch_article_metadata(url, driver=None):
    try:
        response = requests.get(url, headers=HEADERS, timeout=5)
        soup = BeautifulSoup(response.text, "html.parser")
        result = _parse_article_title_and_date(soup)
        if all(result):
            return result
    except Exception:
        pass

    if driver is None:
        return None, None

    try:
        driver.get(url)
        time.sleep(2)
        soup = BeautifulSoup(driver.page_source, "html.parser")
        return _parse_article_title_and_date(soup)
    except Exception:
        return None, None


def collect_news_links(search_url, driver=None):
    selector = (
        "a[href^='https://n.news.naver.com'],"
        "a[href^='https://m.entertain.naver.com']"
    )
    try:
        response = requests.get(search_url, headers=HEADERS, timeout=5)
        soup = BeautifulSoup(response.text, "html.parser")
        links = list({a["href"] for a in soup.select(selector) if a.get("href")})
        if len(links) >= 5:
            return links
    except Exception:
        pass

    if driver is None:
        return []

    try:
        driver.get(search_url)
        time.sleep(2)
        soup = BeautifulSoup(driver.page_source, "html.parser")
        return list({a["href"] for a in soup.select(selector) if a.get("href")})
    except Exception:
        return []


def crawl_naver_articles(
    keywords,
    start_date,
    end_date,
    max_articles_per_keyword=MAX_ARTICLES_PER_KEYWORD,
    headless=HEADLESS,
):
    driver = _build_driver(headless=headless)
    rows = []

    try:
        current_date = _str_to_date(start_date)
        last_date = _str_to_date(end_date)

        while current_date <= last_date:
            day_string = current_date.strftime("%Y.%m.%d")
            for keyword in keywords:
                visited_links = set()
                collected = 0
                fail_count = 0

                for start_index in range(SEARCH_PAGE_START, SEARCH_PAGE_END, SEARCH_PAGE_STEP):
                    if collected >= max_articles_per_keyword:
                        break

                    search_url = (
                        "https://search.naver.com/search.naver"
                        f"?where=news&query={quote(keyword)}"
                        f"&sm=tab_opt&sort=2&photo=0&field=0"
                        f"&pd=3&ds={day_string}&de={day_string}&start={start_index}"
                    )
                    links = collect_news_links(search_url, driver=driver)

                    if not links:
                        fail_count += 1
                        if fail_count >= MAX_CONSECUTIVE_EMPTY_PAGES:
                            break
                        continue

                    fail_count = 0
                    for url in links:
                        if url in visited_links or collected >= max_articles_per_keyword:
                            continue
                        visited_links.add(url)

                        title, published = fetch_article_metadata(url, driver=driver)
                        if title and published:
                            rows.append(
                                {
                                    "keyword": keyword,
                                    "date": published,
                                    "title": title,
                                    "url": url,
                                }
                            )
                            collected += 1
                    time.sleep(1)

            current_date += timedelta(days=1)
    finally:
        driver.quit()

    return pd.DataFrame(rows)


def save_articles_to_excel(df, excel_path):
    if os.path.exists(excel_path):
        workbook = load_workbook(excel_path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "NewsArticles"
        sheet.append(["Keyword", "Title", "Date", "URL"])

    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 60
    sheet.column_dimensions["C"].width = 30
    sheet.column_dimensions["D"].width = 70

    for row in df.itertuples(index=False):
        sheet.append([row.keyword, row.title, row.date, row.url])

    workbook.save(excel_path)


def crawl_comments(url, template="default_society", pool="cbox5", page_size=20):
    oid = url.split("article/")[1].split("/")[0]
    aid = url.split("article/")[1].split("/")[1].split("?")[0]
    page = 1
    rows = []
    headers = {"User-Agent": "Mozilla/5.0", "Referer": url}

    while True:
        callback = f"jQuery{int(time.time() * 1000)}"
        api_url = (
            "https://apis.naver.com/commentBox/cbox/web_neo_list_jsonp.json"
            f"?ticket=news&templateId={template}&pool={pool}"
            f"&_callback={callback}&lang=ko&country=KR"
            f"&objectId=news{oid}%2C{aid}"
            f"&pageSize={page_size}&indexSize=10&listType=OBJECT&pageType=more"
            f"&page={page}&sort=FAVORITE"
        )
        response = requests.get(api_url, headers=headers, timeout=10)
        payload = response.text
        json_str = payload[payload.find("(") + 1 : payload.rfind(")")]
        data = json.loads(json_str)

        comment_list = data.get("result", {}).get("commentList", [])
        if not comment_list:
            break

        for comment in comment_list:
            rows.append(
                {
                    "url": url,
                    "comment_author_id": comment.get("maskedUserId"),
                    "comment_published": comment.get("modTime", comment.get("createTime")),
                    "comment_text": html.unescape(comment.get("contents", "")),
                }
            )

        total = data.get("result", {}).get("totalCommentCount", 0)
        if page * page_size >= total:
            break

        page += 1
        time.sleep(0.3)

    return pd.DataFrame(rows)


def crawl_comments_from_urls(urls):
    all_comments = []
    for index, url in enumerate(pd.Series(urls).dropna(), 1):
        print(f"[{index}] crawling comments: {url}")
        try:
            comments = crawl_comments(url)
            if not comments.empty:
                all_comments.append(comments)
        except Exception as exc:
            print("comment crawling error:", exc)
            continue

    if not all_comments:
        return pd.DataFrame(columns=["url", "comment_author_id", "comment_published", "comment_text"])
    return pd.concat(all_comments, ignore_index=True)


def optimized_filter_bot_comments(df):
    filtered_input = df.copy()
    filtered_input["ts"] = pd.to_datetime(filtered_input["comment_published"], errors="coerce")
    filtered_input.dropna(subset=["ts"], inplace=True)
    filtered_input.sort_values(["comment_author_id", "ts"], inplace=True)
    filtered_input["length"] = filtered_input["comment_text"].fillna("").str.len()
    filtered_input["repetition_ratio"] = (
        filtered_input["comment_text"]
        .fillna("")
        .str.split()
        .map(lambda words: 1 - len(set(words)) / len(words) if len(words) > 0 else 0)
    )
    filtered_input["interval"] = (
        filtered_input.groupby("comment_author_id")["ts"].diff().dt.total_seconds()
    )

    stats = filtered_input.groupby("comment_author_id").agg(
        comment_count=("comment_text", "size"),
        time_var_sec=(
            "interval",
            lambda values: np.var(values.dropna()) if len(values.dropna()) >= 2 else np.nan,
        ),
        length_var=("length", "var"),
        avg_rep_ratio=("repetition_ratio", "mean"),
    )
    stats["time_var_hr"] = stats["time_var_sec"] / 3600

    def ai_prob_series(texts):
        series = texts.fillna("").astype(str)
        sentence_lengths = series.str.split(".").map(
            lambda segments: [len(segment) for segment in segments if segment]
        )
        var_len = sentence_lengths.map(lambda lengths: np.var(lengths) if len(lengths) > 1 else np.nan)
        punct_count = series.str.count(r"[.!?]")
        sentence_count = series.str.count(r"\.") + 1
        punct_ratio = punct_count / sentence_count
        emotional = series.str.contains(r"ㅋㅋ|ㅎㅎ|ㅠㅠ|!!!|\?\?\?")
        return (
            (var_len < 50).fillna(0).astype(float)
            + (punct_ratio > 0.8).astype(float)
            + (~emotional & (series.str.len() > 50)).astype(float)
        ) / 3.0

    filtered_input["ai_prob"] = ai_prob_series(filtered_input["comment_text"])
    ai_stats = filtered_input.groupby("comment_author_id")["ai_prob"].mean().rename("avg_ai_prob")
    stats = stats.join(ai_stats, how="left").fillna(0)

    stats["cond_speed"] = ((stats["comment_count"] >= 5) & (stats["time_var_hr"] < 0.2)).astype(int)
    stats["cond_ai"] = (stats["avg_ai_prob"] > 0.5).astype(int)
    stats["suspicious_score"] = (
        (stats["length_var"] < 200).astype(float) * 0.2
        + (stats["avg_rep_ratio"] > 0.6).astype(float) * 0.3
    )
    stats["cond_suspicious"] = (stats["suspicious_score"] > 0.3).astype(int)
    stats["high_rep"] = (stats["avg_rep_ratio"] > 0.6).astype(int)
    stats["low_len_var"] = (stats["length_var"] < 200).astype(int)

    # Keep the threshold from the old Final_crawling module.
    stats["is_bot"] = (
        stats["cond_speed"]
        + stats["cond_ai"]
        + stats["cond_suspicious"]
        + stats["high_rep"]
        + stats["low_len_var"]
    ) >= 2

    bot_users = stats.index[stats["is_bot"]].tolist()
    filtered_df = filtered_input[~filtered_input["comment_author_id"].isin(bot_users)].drop(
        ["ts", "length", "repetition_ratio", "interval", "ai_prob"],
        axis=1,
    )
    return filtered_df.reset_index(drop=True), stats.reset_index()


def run_pipeline():
    articles_df = crawl_naver_articles(
        keywords=KEYWORDS,
        start_date=START_DATE,
        end_date=END_DATE,
        max_articles_per_keyword=MAX_ARTICLES_PER_KEYWORD,
        headless=HEADLESS,
    )
    save_articles_to_excel(articles_df, ARTICLES_XLSX)

    comments_df = crawl_comments_from_urls(articles_df["url"] if "url" in articles_df else [])
    comments_df.to_csv(COMMENTS_RAW_CSV, index=False, encoding="utf-8-sig")

    filtered_df, stats_df = optimized_filter_bot_comments(comments_df)
    filtered_df.to_csv(COMMENTS_FILTERED_CSV, index=False, encoding="utf-8-sig")
    stats_df.to_csv(BOT_STATS_CSV, index=False, encoding="utf-8-sig")

    print("Saved:", ARTICLES_XLSX, COMMENTS_RAW_CSV, COMMENTS_FILTERED_CSV, BOT_STATS_CSV)


if __name__ == "__main__":
    run_pipeline()
